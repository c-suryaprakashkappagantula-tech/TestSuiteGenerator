# -*- coding: utf-8 -*-
"""demo_doc_generator.py - 'Test Case Presenter Guide' Word document.

Generated with every suite. A generic, OFFLINE, fully deterministic (NO LLM)
client-facing narration guide for a screen-share walkthrough of ANY feature.

Structure (mirrors the hand-edited sample, MWTGPROV-4074):
  - Title block: feature title + "Test Case Presenter Guide" subtitle + tagline + "<FID> | <PI>"
  - Opening narrative (shaded single-cell table)
  - "How to use this guide" (Heading 1) + 3 bullets
  - Per test case: Heading 2 "TCnn | <label>", a "What to say" plain-English
    narration paragraph (first-person-plural, what + why, no jargon), a shaded
    single-cell "Key message:" table, and conditional "Presenter caution:" /
    "Scope note:" asides derived from grounding (never hardcoded)
  - "Closing summary" (Heading 1) + one quoted paragraph
  - "Quick mental model" (Heading 2): TCs grouped by theme, non-empty buckets only

Everything is composed deterministically from each test case's OWN cleaned label
and one-liner (shared with the Excel High Level Scenarios sheet) plus per-category
sentence templates. A banned-phrase safety net scrubs any client-facing string.
Read-only on its inputs; never raises into the caller. No live Jira/QMetry calls.
"""
from __future__ import annotations

import re
from collections import OrderedDict
from datetime import datetime
from pathlib import Path
from typing import Callable, List, Optional

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .client_text import _HLS_BANNED, _hls_description, _hls_scenario_name
from .config import OUTPUTS
from .text_utils import as_text  # MCP may hand back {'text': ...} dicts for long fields

NAVY = '0B2A4A'
TEAL = '0E9AA7'
OPEN_FILL = 'EEF3F8'   # light blue — Opening narrative table
KEY_FILL = 'FFF4D6'    # light amber — Key message table

LQUO = '\u201c'   # “
RQUO = '\u201d'   # ”

# Systems/terms that, when they appear in a TC's source steps but NOT in the
# feature objective, signal unrelated contamination worth a presenter caution.
_OFFTOPIC_TERMS = (
    'prr', 'amdocs', 'mediation', 'nbop_mig', 'ne portal', 'trigger api',
    'service grouping', 'century report', 'transaction_history',
)

# A one-liner that carries any of these is too technical for client narration
# (snake_case/ALL-CAPS identifiers, table/field/file names) — fall back to the
# plain cleaned label instead of echoing it verbatim.
_JARGON_RE = re.compile(
    r'(?:\b[A-Z][A-Z0-9]{2,}(?:_[A-Z0-9]+)+\b'      # STG_TMO_WEARABLE_RECON, TMO_PAIRING_ID
    r'|\b\w+_\w+\b'                                   # any snake_case token
    r'|\.[a-z]{2,4}\b'                               # .csv/.xml file extensions
    r'|\bhttp[s]?://'                                 # URLs
    r'|\bENV_IDENTIFIER\b|\bROUTEDTO\b|\bMDN\b)'      # bare uppercase identifiers
)
# Cues that a sentence is declaring something out of scope.
_SCOPE_CUES = (
    'out of scope', 'not in scope', 'excluded from scope', 'excluded',
    'cross-environment fan-out', 'fan-out',
)


# ------------------------------------------------------------------ text helpers
def _clean(t: str) -> str:
    try:
        from .output_sanitizer import clean_text
        t = clean_text(t or '') or ''
    except Exception:
        t = t or ''
    return t.strip()


def _strip_panels(text: str) -> str:
    return re.sub(r'\{panel[^}]*\}.*?\{panel\}', ' ', as_text(text), flags=re.S | re.I)


def _plain(text: str) -> str:
    """Collapse markup/whitespace to a single clean plain-English string."""
    s = _clean(_strip_panels(as_text(text)))
    s = re.sub(r'\{[^}]*\}', ' ', s)                 # stray macros
    s = re.sub(r'https?://\S+', ' ', s)              # URLs
    s = re.sub(r'[*_#|]+', ' ', s)                   # wiki markup
    s = re.sub(r'\s+', ' ', s).strip()
    return s


def _first_sentences(text: str, n: int = 3) -> str:
    s = _plain(text)
    if not s:
        return ''
    parts = re.split(r'(?<=[.!?])\s+', s)
    out = ' '.join(parts[:n]).strip()
    return out


def _scrub(text: str, fallback: str) -> str:
    """Banned-phrase safety net: if any banned phrase surfaces, use the fallback."""
    low = (text or '').lower()
    if any(b in low for b in _HLS_BANNED):
        return fallback
    return text or fallback


# ------------------------------------------------------------------ docx helpers
def _shade(cell, hex_fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tcPr.append(shd)


def _run(p, text, bold=False, size=None, color=None):
    r = p.add_run(text)
    r.bold = bold
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return r


def _shaded_cell_table(doc, text: str, fill: str):
    """A 1x1 'Normal Table' whose single cell is shaded `fill`."""
    t = doc.add_table(rows=1, cols=1)
    try:
        t.style = 'Normal Table'
    except Exception:
        pass
    cell = t.rows[0].cells[0]
    cell.text = text
    _shade(cell, fill)
    return t


def _quote(text: str) -> str:
    return '%s%s%s' % (LQUO, text.strip(), RQUO)


# ------------------------------------------------------------------ composition
def _feature_objective(title: str, blob: str) -> str:
    """Short, clean noun phrase describing what the feature is about (for caution text).

    Strips bracketed tags like '[NSLNM, INTG]:' and 'New MVNO -' prefixes, drops a
    trailing jargon parenthetical/file token, and keeps it readable."""
    t = _plain(title) or 'this feature'
    t = re.sub(r'^\s*\[[^\]]*\]\s*:?\s*', '', t)             # drop leading [TAG]:
    t = re.sub(r'^\s*(new\s+mvno|intg|nslnm)\s*[-:]\s*', '', t, flags=re.I)
    t = re.sub(r'\s*\([^)]*\)\s*$', '', t)                   # drop trailing (CHRTR_deltas_...)
    t = re.sub(r'\s*(feature|story|epic|process)\s*$', '', t, flags=re.I).strip()
    t = re.sub(r'\s+', ' ', t).strip(' -:')
    if not t:
        return 'this feature'
    return (t[:1].lower() + t[1:])[:70]


def _tc_label(tc, idx: int) -> str:
    """Short human label; if it carries a banned phrase, fall back to a neutral label
    so no banned term reaches the client-facing heading."""
    label = _hls_scenario_name(tc, idx)
    low = (label or '').lower()
    if any(b in low for b in _HLS_BANNED):
        return 'Scenario %02d' % idx
    return label


def _subject(tc, label: str) -> str:
    """The plain-English 'what' clause for narration.

    Starts from the shared one-liner but REJECTS it if it carries technical jargon
    (snake_case/ALL-CAPS identifiers, file/table/field names, off-topic systems) or
    is empty — in those cases it falls back to the cleaned human label so the
    narration stays client-facing. Never echoes raw step text verbatim."""
    one = _hls_description(tc, name=label) or ''
    one = re.sub(r'(?i)^\s*(verify|validate|ensure|confirm|check)\s+(that\s+)?', '', one).strip()
    one = one.rstrip('.;, ').strip()
    low = one.lower()
    jargon = bool(_JARGON_RE.search(one)) or any(t in low for t in _OFFTOPIC_TERMS)
    if not one or jargon:
        one = (label or 'the scenario').strip()
    return _mid_sentence(one)


def _mid_sentence(phrase: str) -> str:
    """Make a phrase read naturally in the middle of a sentence: lowercase each word
    unless it is an ALL-CAPS acronym (SFTP, UAT, MDN) or a dotted/branded token."""
    words = []
    for w in (phrase or '').split(' '):
        if not w:
            continue
        core = re.sub(r'[^A-Za-z0-9]', '', w)
        if core.isupper() and len(core) > 1:      # acronym — keep as-is
            words.append(w)
        elif w[:1].isupper() and w[1:] != w[1:].lower():  # CamelCase/branded — keep
            words.append(w)
        else:
            words.append(w[:1].lower() + w[1:])
    return ' '.join(words)


def _category(tc) -> str:
    return (getattr(tc, 'category', '') or 'Happy Path').strip()


def _tc_text_blob(tc) -> str:
    """All of a TC's own text (summary/desc/steps) for keyword cues + off-topic scan."""
    parts = [getattr(tc, 'summary', '') or '', getattr(tc, 'description', '') or '']
    for s in (getattr(tc, 'steps', None) or []):
        parts.append(getattr(s, 'summary', '') or '')
        parts.append(getattr(s, 'expected', '') or '')
    return _plain(' '.join(parts)).lower()


def _has(blob: str, *cues: str) -> bool:
    return any(c in blob for c in cues)


def _narration_template(category: str, blob: str) -> str:
    """Pick a per-category narration frame refined by keyword cues. Frame has one
    '{subject}' slot. Deterministic selection: refine by cue first, else category."""
    # visibility / reporting — only genuine result-visibility intent, not the mere
    # presence of the words "report"/"audit" (which name the feature, not the check).
    if _has(blob, 'is visible', 'visibility', 'clearly see', 'outcome is', 'result is visible',
            'monitor', 'monitoring', 'alert', 'dashboard shows', 'processing-result'):
        return ('This scenario checks that {subject} is clearly visible after processing, '
                'so support can see what happened without manual investigation.')
    # negative / exception
    if category in ('Negative', 'Edge Case') or _has(
            blob, 'invalid', 'blank', 'empty', 'missing', 'not found', 'fail', 'incorrect',
            'exception', ' bad ', 'does not arrive', 'cannot be found'):
        return ('In this scenario, we deliberately test {subject}. The system should handle it '
                'safely - one bad record should not stop the rest, and the problem should be '
                'visible rather than causing a silent failure.')
    # regression / no-impact
    if category == 'Regression' or _has(blob, 'not impacted', 'no impact', 'existing', 'unchanged'):
        return ('This scenario confirms that {subject} does not affect existing functionality. '
                'Anything that was already working should keep working exactly as before.')
    # e2e
    if category == 'E2E' or _has(blob, 'end to end', 'end-to-end', 'full flow'):
        return ('This scenario walks the full flow end to end for {subject}, confirming the '
                'pieces work together and the final result is correct.')
    # validation / compare
    if _has(blob, 'validate', 'verify', 'compare', 'reconcile', 'match', 'pairing', 'mismatch'):
        return ('This is where the main checking happens. We compare {subject} and confirm the '
                'system reacts correctly - correcting what is wrong and leaving what is already '
                'right unchanged.')
    # happy-path default
    return ('In this scenario, we are checking that {subject}. The goal is to confirm the '
            'expected outcome happens cleanly and the process completes as designed.')


def _narration(tc, label: str) -> str:
    cat = _category(tc)
    blob = _tc_text_blob(tc)
    subject = _subject(tc, label)
    frame = _narration_template(cat, blob)
    text = frame.format(subject=subject)
    fallback = ('In this scenario, we are checking that %s behaves as the feature requires.'
                % (label.lower() if label else 'this scenario'))
    return _scrub(text, fallback)


def _key_message(tc, label: str) -> str:
    cat = _category(tc)
    blob = _tc_text_blob(tc)
    # Category-specific closers read better where they clearly apply.
    if cat in ('Negative', 'Edge Case') or _has(blob, 'invalid', 'missing', 'not found', 'blank', 'empty'):
        msg = 'One bad record must not stop the valid ones.'
    elif cat == 'Regression' or _has(blob, 'no impact', 'not impacted', 'existing'):
        msg = 'The new feature must not break existing processing.'
    elif _has(blob, 'already', 'unchanged', 'matches'):
        msg = 'If the record is already correct, leave it unchanged.'
    elif _has(blob, 'is visible', 'visibility', 'monitor', 'monitoring', 'processing-result', 'outcome is'):
        msg = 'Processing is not complete until the outcome is visible and understandable.'
    else:
        subject = _subject(tc, label)
        msg = '%s completes correctly and safely.' % (subject[:1].upper() + subject[1:] if subject else 'The scenario')
    fallback = 'The scenario behaves as the feature requires.'
    return _scrub(msg, fallback)


def _tc_step_text(tc) -> str:
    """Just the TC's step summaries + expected (for the grounding check)."""
    parts = []
    for s in (getattr(tc, 'steps', None) or []):
        parts.append(getattr(s, 'summary', '') or '')
        parts.append(getattr(s, 'expected', '') or '')
    return ' '.join(parts).strip()


def _caution_note(tc, objective: str, evidence_blob: str) -> str:
    """Emit a caution only when a TC's source steps reference a concept that is NOT
    grounded in the feature objective — i.e. the step text is banned-and-ungrounded
    against the feature evidence blob, OR it names an off-topic system/term that does
    not appear anywhere in the feature evidence. Derived, never hardcoded. Fail open."""
    step_text = _tc_step_text(tc)
    if not step_text:
        return ''
    blob_low = (evidence_blob or '').lower()
    # An off-topic term the steps mention that is absent from the feature evidence.
    detected = next((t for t in _OFFTOPIC_TERMS
                     if t in step_text.lower() and t not in blob_low), '')
    banned_ungrounded = False
    if not detected:
        try:
            from .qmetry_pattern_library import _step_is_banned_and_ungrounded
            banned_ungrounded = bool(_step_is_banned_and_ungrounded(step_text, evidence_blob or ''))
        except Exception:
            banned_ungrounded = False
    if not detected and not banned_ungrounded:
        return ''
    # Never surface a banned phrase as the named "system" — genericize it.
    if detected and any(b in detected.lower() for b in _HLS_BANNED):
        detected = ''
    system = (detected.upper() if detected and len(detected) <= 5
              else (detected or 'unrelated'))
    return ('Keep the narration aligned to the %s objective. The detailed source steps appear to '
            'reference unrelated %s processing and should be reviewed separately.'
            % (objective, system))


def _scope_note(tc) -> str:
    """Emit a scope note when the TC text declares something out of scope."""
    blob = _tc_text_blob(tc)
    hit = next((c for c in _SCOPE_CUES if c in blob), '')
    if not hit:
        return ''
    # Try to recover the excluded thing from the sentence that carries the cue.
    full = _plain(' '.join([getattr(tc, 'description', '') or ''] +
                           [getattr(s, 'summary', '') or '' for s in (getattr(tc, 'steps', None) or [])]))
    excluded = hit
    for sent in re.split(r'(?<=[.!?])\s+', full):
        if hit in sent.lower():
            excluded = sent.strip().rstrip('.')[:120] or hit
            break
    return ('Present this as scope-dependent. The source material notes that %s is excluded '
            'from scope.' % excluded)


# ----- quick mental model buckets (fixed order; render only non-empty) ---------
_BUCKETS = [
    ('Receive the file',
     lambda cat, blob: _has(blob, 'receive', 'pick up', 'pickup', 'parse', 'ingest', 'load', 'schedul', 'sftp')),
    ('Validate and compare',
     lambda cat, blob: _has(blob, 'validate', 'verify', 'compare', 'reconcile', 'match', 'pairing')),
    ('Correct the mismatch',
     lambda cat, blob: _has(blob, 'correct', 'update', ' fix', 'mismatch', 'discrepanc')),
    ('Handle exceptions safely',
     lambda cat, blob: cat in ('Negative', 'Edge Case') or _has(
         blob, 'blank', 'empty', 'invalid', 'missing', 'not found', 'fail', 'exception', ' bad ')),
    ('Avoid unnecessary updates',
     lambda cat, blob: _has(blob, 'already', 'matches', 'unchanged', 'no update', 'outside scope', 'skip')),
    ('Regression - no impact',
     lambda cat, blob: cat == 'Regression' or _has(blob, 'existing', 'not impacted', 'no impact')),
    ('Show the outcome',
     lambda cat, blob: _has(blob, 'is visible', 'visibility', 'monitor', 'monitoring',
                            'processing-result', 'outcome is', 'result is visible')),
]


def _mental_model(tcs) -> "OrderedDict[str, List[str]]":
    out: "OrderedDict[str, List[str]]" = OrderedDict((label, []) for label, _ in _BUCKETS)
    for i, tc in enumerate(tcs, 1):
        cat = _category(tc)
        blob = _tc_text_blob(tc)
        tag = 'TC%02d' % i
        for label, rule in _BUCKETS:
            try:
                if rule(cat, blob):
                    out[label].append(tag)
            except Exception:
                pass
    return OrderedDict((k, v) for k, v in out.items() if v)


def _oxford(tcs: List[str]) -> str:
    """'TC01', 'TC01 and TC02', 'TC01, TC02 and TC03'."""
    if not tcs:
        return ''
    if len(tcs) == 1:
        return tcs[0]
    if len(tcs) == 2:
        return '%s and %s' % (tcs[0], tcs[1])
    return '%s and %s' % (', '.join(tcs[:-1]), tcs[-1])


def _opening_narrative(title: str, desc: str, objective: str) -> str:
    """One plain paragraph. Built from the feature description with the (repeated)
    title prefix removed and technical jargon sentences dropped, so it reads as a
    clean client-facing restatement rather than echoing internal field names."""
    body = _plain(desc)
    plain_title = _plain(title)
    # Descriptions typically begin by repeating the title — strip that prefix.
    if plain_title and body.lower().startswith(plain_title.lower()):
        body = body[len(plain_title):].lstrip(' -|:').strip()
    # Keep only the opening sentences that are free of technical jargon.
    clean_sents = []
    for sent in re.split(r'(?<=[.!?])\s+', body):
        s = sent.strip()
        if not s:
            continue
        if _JARGON_RE.search(s) or any(t in s.lower() for t in _OFFTOPIC_TERMS):
            break  # stop at the first technical sentence
        clean_sents.append(s)
        if len(clean_sents) >= 3:
            break
    gist = ' '.join(clean_sents).strip()
    if gist:
        text = 'The overall concept is simple. %s' % gist
    else:
        obj = objective[:1].upper() + objective[1:] if objective else 'This feature'
        text = ('The overall concept is simple. %s receives its source data, validates it, and '
                'acts on it so the right outcome is reached safely.' % obj)
    fallback = ('The overall concept is simple. We receive the source data, validate it, and act '
                'on it so the right outcome is reached safely.')
    return _scrub(text, fallback)


# Map each bucket label to a short theme phrase for the closing summary.
_THEME_PHRASE = {
    'Receive the file': 'the source data can be received and read',
    'Validate and compare': 'the data can be validated and compared',
    'Correct the mismatch': 'incorrect information can be corrected',
    'Handle exceptions safely': 'problems are handled safely without stopping the valid work',
    'Avoid unnecessary updates': 'records that are already correct are left unchanged',
    'Regression - no impact': 'existing functionality keeps working',
    'Show the outcome': 'the outcome is clearly visible afterwards',
}


def _closing_summary(buckets: "OrderedDict[str, List[str]]") -> str:
    themes = [_THEME_PHRASE[b] for b in buckets if b in _THEME_PHRASE]
    if themes:
        body = _oxford(themes) if len(themes) <= 2 else '%s, and %s' % (', '.join(themes[:-1]), themes[-1])
        text = 'Across these scenarios, we are proving that %s.' % body
    else:
        text = ('Across these scenarios, we are proving that the feature behaves correctly and '
                'safely across the situations it must handle.')
    fallback = ('Across these scenarios, we are proving that the feature behaves correctly and '
                'safely across the situations it must handle.')
    return _scrub(text, fallback)


# ------------------------------------------------------------------ main
def generate_demo_doc(suite, jira=None, chalk=None, log: Callable = print,
                      out_dir: Optional[Path] = None) -> Optional[Path]:
    """Write <FID>_<PI>_Test_Case_Presenter_Guide_<ts>.docx. Returns path (None on failure)."""
    try:
        return _build(suite, jira, chalk, log, out_dir)
    except Exception as e:  # a presenter guide must never break suite generation
        if log:
            log('[DEMO-DOC] skipped: %s' % str(e)[:120])
        return None


def _build(suite, jira, chalk, log, out_dir):
    fid = getattr(suite, 'feature_id', '') or (getattr(jira, 'key', '') if jira else '')
    title = (getattr(suite, 'feature_title', '') or (getattr(jira, 'summary', '') if jira else '')
             or fid or 'Feature')
    pi = getattr(suite, 'pi', '') or (getattr(jira, 'pi', '') if jira else '') or ''
    tcs = list(getattr(suite, 'test_cases', None) or [])

    # Feature objective blob — Chalk/Jira desc + scope + AC — used for opening + grounding.
    desc = _plain(getattr(suite, 'feature_desc', '') or (as_text(getattr(jira, 'description', '')) if jira else ''))
    scope = _plain(getattr(suite, 'scope', '') or (getattr(chalk, 'scope', '') if chalk is not None else ''))
    objective = _feature_objective(title, ' '.join(x for x in (title, desc, scope) if x))

    # Feature evidence blob drives the grounded caution check (fail open on any error).
    try:
        from .step_assembler import _feature_evidence_blob
        evidence_blob = _feature_evidence_blob(suite) or ''
    except Exception:
        evidence_blob = ' '.join(x for x in (title, desc, scope) if x)

    doc = Document()
    sec = doc.sections[0]
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Inches(0.79))
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)

    # ---------------- title block
    p = doc.add_paragraph()
    p.style = doc.styles['Title']
    _run(p, title)
    _run(p, '\nTest Case Presenter Guide')
    _run(doc.add_paragraph(), 'Plain-English narration for screen-share walkthroughs', bold=True)
    _run(doc.add_paragraph(), '%s | %s' % (fid, pi or 'PI'))

    # ---------------- opening narrative (shaded)
    _shaded_cell_table(doc, 'Opening narrative\n%s'
                       % _quote(_opening_narrative(title, desc or scope, objective)), OPEN_FILL)

    # ---------------- how to use
    doc.add_heading('How to use this guide', level=1)
    for b in (
        'Keep the test case open on screen and use the %sWhat to say%s section as your narration.' % (LQUO, RQUO),
        'Use the %sKey message%s as a short closing statement before moving to the next test case.' % (LQUO, RQUO),
        'The wording is intentionally non-technical so business, operations, and technical '
        'audiences can follow the same walkthrough.',
    ):
        doc.add_paragraph(b, style='List Bullet')
    doc.add_paragraph('')

    # ---------------- per TC
    for i, tc in enumerate(tcs, 1):
        label = _tc_label(tc, i)
        doc.add_heading('TC%02d | %s' % (i, label), level=2)

        p = doc.add_paragraph()
        _run(p, 'What to say', bold=True)
        _run(p, '\n', bold=True)
        _run(p, _quote(_narration(tc, label)))

        _shaded_cell_table(doc, 'Key message: %s' % _key_message(tc, label), KEY_FILL)

        caution = _caution_note(tc, objective, evidence_blob)
        if caution:
            cp = doc.add_paragraph()
            _run(cp, 'Presenter caution: ', bold=True)
            _run(cp, caution)
        snote = _scope_note(tc)
        if snote:
            sp = doc.add_paragraph()
            _run(sp, 'Scope note: ', bold=True)
            _run(sp, snote)

        doc.add_paragraph('')

    # ---------------- closing summary
    buckets = _mental_model(tcs)
    doc.add_heading('Closing summary', level=1)
    doc.add_paragraph(_quote(_closing_summary(buckets)))

    # ---------------- quick mental model
    doc.add_heading('Quick mental model', level=2)
    for label, tags in buckets.items():
        mp = doc.add_paragraph()
        _run(mp, '%s: ' % label, bold=True)
        _run(mp, _oxford(tags))

    out_dir = Path(out_dir) if out_dir else OUTPUTS
    out_dir.mkdir(parents=True, exist_ok=True)
    name = '%s_%s_Test_Case_Presenter_Guide_%s.docx' % (
        fid, (pi or 'PI').replace(' ', ''), datetime.now().strftime('%Y%m%d_%H%M%S'))
    path = out_dir / name
    doc.save(str(path))
    if log:
        log('[DEMO-DOC] Test Case Presenter Guide saved: %s' % name)
    return path
