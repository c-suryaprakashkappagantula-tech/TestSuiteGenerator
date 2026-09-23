"""
mine_qmetry_patterns.py — Regenerate modules/qmetry_pattern_library.json.

Mines the human "house style" from exported QMetry manual test-suite .xlsx files
(the QMetry TestCases export layout: one header row, TCs identified by a non-empty
"Issue Key" cell, step rows continue with a blank Issue Key). Produces the JSON
consumed by modules/qmetry_pattern_library.py.

Usage:
    python tools/mine_qmetry_patterns.py <file1.xlsx> <file2.xlsx> ...

Deterministic: output is frequency-ordered, no randomness.
"""
import json
import re
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import openpyxl

OUT = Path(__file__).resolve().parents[1] / 'modules' / 'qmetry_pattern_library.json'

# QMetry export column indices (1-based)
C_KEY, C_SUMMARY, C_DESC, C_PRECOND = 1, 2, 3, 4
C_STEP, C_EXPECTED = 17, 19

_PRECON_KEYWORDS = [
    'activate', 'port', 'wearable', 'sim', 'feature', 'line status', 'suspend',
    'restore', 'swap', 'mdn', 'hotline', 'sync', 'inquiry', 'plan', 'device',
    'account', 'reset', 'admin', 'permission', 'host',
]
_PRECON_NOISE = re.compile(
    r'software|version|26\.|websheet|previous day|sftp|hld|reference number', re.I)


def _norm(s):
    s = str(s or '').replace('\r', ' ')
    s = re.sub(r'^\s*(?:step\s*)?\d+[\.\)]\s*', '', s, flags=re.I)
    return re.sub(r'\s+', ' ', s).strip()


def _load_tcs(files):
    tcs = []
    for f in files:
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb['TestCases'] if 'TestCases' in wb.sheetnames else wb.worksheets[0]
        cur = None
        for r in range(2, ws.max_row + 1):
            key = ws.cell(r, C_KEY).value
            if key:
                cur = {
                    'summary': str(ws.cell(r, C_SUMMARY).value or ''),
                    'desc': str(ws.cell(r, C_DESC).value or ''),
                    'precond': str(ws.cell(r, C_PRECOND).value or ''),
                    'steps': [],
                }
                tcs.append(cur)
            if cur is not None:
                ss = ws.cell(r, C_STEP).value
                ex = ws.cell(r, C_EXPECTED).value
                if (ss and str(ss).strip()) or (ex and str(ex).strip()):
                    cur['steps'].append((str(ss or ''), str(ex or '')))
    return tcs


def _top_steps(tcs, pred, n=4):
    c = Counter()
    for t in tcs:
        for ss, ex in t['steps']:
            s, e = _norm(ss), _norm(ex)
            if pred(s, e) and 8 < len(s) < 160:
                c[(s[:150], e[:150])] += 1
    return [{'summary': k[0], 'expected': k[1], 'freq': v} for k, v in c.most_common(n)]


def mine(files):
    tcs = _load_tcs(files)
    alltext = '\n'.join(ss + '\n' + ex for t in tcs for ss, ex in t['steps'])

    # Preconditions
    pc = Counter()
    for t in tcs:
        for ln in t['precond'].split('\n'):
            ln = _norm(ln)
            if 10 < len(ln) < 110:
                pc[ln] += 1
    by_kw = defaultdict(list)
    for p, _n in pc.most_common():
        pl = p.lower()
        for kw in _PRECON_KEYWORDS:
            if re.search(r'\b' + re.escape(kw), pl) and len(by_kw[kw]) < 6:
                by_kw[kw].append(p)
    safe = [p for p, _n in pc.most_common(40) if not _PRECON_NOISE.search(p)][:12]

    # DB tables
    tables = Counter(re.findall(r'(NBOP_MIG_[A-Z_]+)', alltext))
    canonical = [{'table': tb, 'freq': n} for tb, n in tables.most_common()
                 if n >= 30 and len(tb) > 10]

    # Closers + pairing
    closers = {
        'century_txn_trace': _top_steps(
            tcs, lambda s, e: 'century db' in s.lower()
            and ('txn' in s.lower() or 'transaction log' in s.lower())),
        'db_table_validation': _top_steps(
            tcs, lambda s, e: 'nbop_mig_' in (s + e).lower()
            or ('nsl' in s.lower() and 'table' in s.lower())),
        'genesis_portal_verify': _top_steps(
            tcs, lambda s, e: 'genesis' in (s + e).lower()),
        'e2e_flow': _top_steps(
            tcs, lambda s, e: 'end to end' in s.lower() or 'end-to-end' in s.lower()),
    }
    pairing = _top_steps(
        tcs, lambda s, e: 'century report' in s.lower() and s.lower().startswith('validate'), 5)

    # Negative vocab
    neg_fields = Counter(re.findall(
        r'"(responseCode|errorCode|errorDetails|reason|message|statusCode|description|status)"', alltext))
    neg_msgs = Counter(re.findall(
        r'"(Bad Request|Invalid [A-Za-z ]{3,20}|Not Found|Unauthorized|Forbidden|[A-Za-z ]{0,15}not found)"', alltext))
    resp_codes = Counter(re.findall(r'"(?:statusCode|responseCode)"\s*:?\s*"?(\d{3})', alltext))

    # Real API operation names
    cands = Counter(re.findall(
        r'"([A-Z][a-zA-Z]+(?:[A-Z][a-z]+){2,}|[A-Za-z]+(?:\s[A-Z][a-zA-Z]+){1,3})"', alltext))
    trig = Counter(re.findall(r'[Tt]rigger(?:\s+the)?\s+"?([A-Z][A-Za-z]{6,40})"?\s+api', alltext))
    apis = Counter()
    for c, n in list(cands.items()) + list(trig.items()):
        c = c.strip()
        if len(c) >= 10 and n >= 5 and not c.isupper() and not c.islower():
            apis[c] += n

    return {
        'meta': {
            'generated': str(date.today()),
            'source': 'QMetry manual suites (MDA/NSL provisioning)',
            'source_files': [Path(f).name for f in files],
            'total_tcs': len(tcs),
            'total_steps': sum(len(t['steps']) for t in tcs),
        },
        'preconditions': {
            'global_top': [{'text': p, 'freq': n} for p, n in pc.most_common(30)],
            'by_keyword': {k: v for k, v in by_kw.items() if v},
            'safe_defaults': safe,
        },
        'db_tables': {'canonical': canonical},
        'closers': closers,
        'century_pairing_exemplars': pairing,
        'negative_vocab': {
            'error_fields': [f for f, _ in neg_fields.most_common()],
            'error_messages': [m for m, n in neg_msgs.most_common(15) if n >= 3],
            'response_codes': [c for c, _ in resp_codes.most_common(10)],
        },
        'api_names': [a for a, _ in apis.most_common(24)],
    }


def main(argv):
    if not argv:
        print(__doc__)
        return 1
    lib = mine(argv)
    OUT.write_text(json.dumps(lib, indent=2, ensure_ascii=False), encoding='utf-8')
    m = lib['meta']
    print('Wrote %s' % OUT)
    print('  %d TCs / %d steps from %d file(s)'
          % (m['total_tcs'], m['total_steps'], len(m['source_files'])))
    print('  canonical tables: %s'
          % ', '.join(t['table'] for t in lib['db_tables']['canonical']))
    return 0


if __name__ == '__main__':
    raise SystemExit(main(sys.argv[1:]))
