"""Compare pre/post crawl-grounding workbooks for the regression set.

Spec: tsg-nbop-crawl-grounded-ui-steps (Req 8.1, 8.2, 8.4, 9.1-9.6).

Usage (cwd TestSuiteGenerator):
    python tools/compare_grounding_regression.py \
        --before outputs/grounding_regression/before \
        --after outputs/grounding_regression/after \
        --features <KEY> <KEY> ... [--typed40 <KEY> ...] [--ui <KEY> ...] \
        --out outputs/grounding_regression/summary.json

Feature keys come only from the CLI or from ``<before>/manifest.json``
(no hardcoded keys). A feature is UI when its manifest classification is
``ui`` or it is listed in ``--ui``. Grounding Report thresholds are applied
to every feature whose after-workbook contains a ``Grounding Report`` sheet.
"""
from __future__ import annotations

import argparse
import json
import os
import sys
from typing import Any, Dict, List, Optional

import openpyxl

TEST_CASES_SHEET = "Test Cases"
CONTRACT_SHEET = "Execution Contract"
REPORT_SHEET = "Grounding Report"

# Req 9 thresholds
GROUNDED_MIN = 55.0
TYPED_MIN_HIGH = 40.0
TYPED_MIN_DEFAULT = 15.0

UNMATCHED_REASONS = {"no_match", "ambiguous", "crawl_unavailable"}

# Metric labels from design §10 -> output keys
_METRIC_KEYS = {
    "ui_steps": "ui_steps",
    "grounded_steps": "grounded_steps",
    "unmatched_steps": "unmatched_steps",
    "filler_steps (post-pass)": "filler_steps",
    "filler_steps": "filler_steps",
    "closer_steps removed": "closers_removed",
    "typed rows": "typed_rows",
    "total rows": "total_rows",
    "grounded share %": "grounded_share",
    "typed share %": "typed_share",
}


# ---------------------------------------------------------------------------
# Workbook helpers
# ---------------------------------------------------------------------------

def _load(path: str):
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def _sheet_rows(wb, name: str) -> Optional[List[tuple]]:
    if name not in wb.sheetnames:
        return None
    rows = [tuple(r) for r in wb[name].iter_rows(values_only=True)]
    # Normalise trailing empty cells / rows so read-only dims don't matter.
    norm = []
    for r in rows:
        r = list(r)
        while r and r[-1] in (None, ""):
            r.pop()
        norm.append(tuple(r))
    while norm and not norm[-1]:
        norm.pop()
    return norm


def _tc_titles(rows: Optional[List[tuple]]) -> List[str]:
    """Ordered TC titles: the `Summary` column of rows with an S.No value."""
    if not rows:
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    try:
        title_col = header.index("Summary")
    except ValueError:
        title_col = 1
    titles = []
    for r in rows[1:]:
        if r and r[0] not in (None, ""):
            titles.append(str(r[title_col]) if len(r) > title_col and r[title_col] is not None else "")
    return titles


def _first_diff(a: Optional[List[tuple]], b: Optional[List[tuple]]) -> Optional[Dict[str, Any]]:
    a = a or []
    b = b or []
    for i in range(max(len(a), len(b))):
        ra = a[i] if i < len(a) else None
        rb = b[i] if i < len(b) else None
        if ra != rb:
            return {"row": i + 1, "before": list(ra) if ra is not None else None,
                    "after": list(rb) if rb is not None else None}
    return None


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def compare_feature(before_xlsx: str, after_xlsx: str, is_ui: bool) -> dict:
    """Req 8.1/8.2: TC count + ordered titles; non-UI sheets cell-identical."""
    wb_b, wb_a = _load(before_xlsx), _load(after_xlsx)
    try:
        tc_b = _sheet_rows(wb_b, TEST_CASES_SHEET)
        tc_a = _sheet_rows(wb_a, TEST_CASES_SHEET)
        titles_b, titles_a = _tc_titles(tc_b), _tc_titles(tc_a)
        result: Dict[str, Any] = {
            "is_ui": bool(is_ui),
            "tc_count_before": len(titles_b),
            "tc_count_after": len(titles_a),
            "tc_count_equal": len(titles_b) == len(titles_a),
            "titles_equal": titles_b == titles_a,
            "has_grounding_report": REPORT_SHEET in wb_a.sheetnames,
        }
        if titles_b != titles_a:
            result["title_diffs"] = [
                {"index": i + 1,
                 "before": titles_b[i] if i < len(titles_b) else None,
                 "after": titles_a[i] if i < len(titles_a) else None}
                for i in range(max(len(titles_b), len(titles_a)))
                if (titles_b[i] if i < len(titles_b) else None)
                != (titles_a[i] if i < len(titles_a) else None)
            ]
        ok = result["tc_count_equal"] and result["titles_equal"]
        if not is_ui:
            result["sheet_names_equal"] = list(wb_b.sheetnames) == list(wb_a.sheetnames)
            if not result["sheet_names_equal"]:
                result["sheet_names_before"] = list(wb_b.sheetnames)
                result["sheet_names_after"] = list(wb_a.sheetnames)
            for key, name in (("test_cases_identical", TEST_CASES_SHEET),
                              ("contract_identical", CONTRACT_SHEET)):
                rb = _sheet_rows(wb_b, name)
                ra = _sheet_rows(wb_a, name)
                same = rb == ra
                result[key] = same
                if not same:
                    result[key.replace("_identical", "_first_diff")] = _first_diff(rb, ra)
            ok = ok and result["sheet_names_equal"] and result["test_cases_identical"] \
                and result["contract_identical"]
        result["passed"] = bool(ok)
        return result
    finally:
        wb_b.close()
        wb_a.close()


def _to_number(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip().rstrip("%").strip()
    try:
        return float(s)
    except ValueError:
        return None


def ui_metrics(after_xlsx: str) -> dict:
    """Read the Grounding Report metrics block and issue rows."""
    wb = _load(after_xlsx)
    try:
        rows = _sheet_rows(wb, REPORT_SHEET)
    finally:
        wb.close()
    if rows is None:
        return {"has_grounding_report": False}

    metrics: Dict[str, Any] = {"has_grounding_report": True}
    i = 0
    # Metrics block: label | value rows until the first blank row.
    while i < len(rows) and rows[i]:
        label = str(rows[i][0]).strip() if rows[i][0] is not None else ""
        value = rows[i][1] if len(rows[i]) > 1 else None
        key = _METRIC_KEYS.get(label.lower())
        if key:
            metrics[key] = _to_number(value)
        elif label.lower() == "feature":
            metrics["feature"] = value
        i += 1
    # Skip blank rows, then the issue table header.
    while i < len(rows) and not rows[i]:
        i += 1
    issues: List[Dict[str, Any]] = []
    if i < len(rows):
        header = [str(h).strip() if h is not None else "" for h in rows[i]]
        reason_col = header.index("Reason") if "Reason" in header else 4
        for r in rows[i + 1:]:
            if not r:
                continue
            row = {header[c] if c < len(header) and header[c] else f"col{c}": r[c]
                   for c in range(len(r))}
            row["_reason"] = str(r[reason_col]).strip() if len(r) > reason_col and r[reason_col] else ""
            issues.append(row)

    unmatched_rows = [x for x in issues if x["_reason"] in UNMATCHED_REASONS]
    metrics["issue_rows"] = len(issues)
    metrics["unmatched_issue_rows"] = len(unmatched_rows)
    metrics["no_match_rows"] = [
        {k: v for k, v in x.items() if not k.startswith("_")}
        for x in issues if x["_reason"] == "no_match"
    ]
    metrics["issues_match_unmatched"] = (
        metrics.get("unmatched_steps") is not None
        and int(metrics["unmatched_steps"]) == len(unmatched_rows)
    )
    ui, g, u = metrics.get("ui_steps"), metrics.get("grounded_steps"), metrics.get("unmatched_steps")
    metrics["all_ui_classified"] = None not in (ui, g, u) and (g + u) == ui
    return metrics


def check_thresholds(metrics: dict, typed40: bool) -> dict:
    """Apply Req 9 thresholds to a ``ui_metrics`` result."""
    if not metrics.get("has_grounding_report"):
        return {"passed": False, "failures": ["missing Grounding Report sheet"]}
    typed_min = TYPED_MIN_HIGH if typed40 else TYPED_MIN_DEFAULT
    failures = []
    filler = metrics.get("filler_steps")
    if filler is None or filler != 0:
        failures.append(f"filler_steps={filler} (expected 0)")
    if not metrics.get("all_ui_classified"):
        failures.append("grounded + unmatched != ui_steps")
    gs = metrics.get("grounded_share")
    if gs is None or gs < GROUNDED_MIN:
        failures.append(f"grounded_share={gs} < {GROUNDED_MIN}")
    ts = metrics.get("typed_share")
    if ts is None or ts < typed_min:
        failures.append(f"typed_share={ts} < {typed_min}")
    if not metrics.get("issues_match_unmatched"):
        failures.append(
            f"unmatched issue rows={metrics.get('unmatched_issue_rows')} "
            f"!= Unmatched_Steps={metrics.get('unmatched_steps')}")
    return {"passed": not failures, "typed_min": typed_min, "failures": failures}


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def _load_manifest(before_dir: str) -> List[dict]:
    path = os.path.join(before_dir, "manifest.json")
    if not os.path.isfile(path):
        return []
    with open(path, encoding="utf-8") as fh:
        data = json.load(fh)
    return data if isinstance(data, list) else []


def run(before_dir: str, after_dir: str, features: Optional[List[str]],
        typed40: List[str], ui_override: List[str]) -> dict:
    manifest = {m.get("feature"): m for m in _load_manifest(before_dir) if m.get("feature")}
    feats = list(features) if features else list(manifest.keys())
    summary: Dict[str, Any] = {"features": {}, "passed": True}
    for feat in feats:
        entry: Dict[str, Any] = {}
        classification = (manifest.get(feat) or {}).get("classification", "unknown")
        is_ui = classification == "ui" or feat in ui_override
        entry["classification"] = classification
        before = os.path.join(before_dir, f"{feat}.xlsx")
        after = os.path.join(after_dir, f"{feat}.xlsx")
        missing = [p for p in (before, after) if not os.path.isfile(p)]
        if missing:
            entry.update(passed=False, error=f"missing workbook(s): {missing}")
        else:
            cmp = compare_feature(before, after, is_ui)
            entry["comparison"] = cmp
            ok = cmp["passed"]
            if cmp["has_grounding_report"] or is_ui:
                m = ui_metrics(after)
                th = check_thresholds(m, feat in typed40)
                entry["metrics"] = m
                entry["thresholds"] = th
                ok = ok and th["passed"]
            entry["passed"] = ok
        summary["features"][feat] = entry
        summary["passed"] = summary["passed"] and entry["passed"]
    return summary


def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--before", required=True)
    ap.add_argument("--after", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--features", nargs="*", default=None,
                    help="feature keys (default: all manifest entries)")
    ap.add_argument("--typed40", nargs="*", default=[],
                    help="features that need typed share >= 40%%")
    ap.add_argument("--ui", nargs="*", default=[],
                    help="treat these features as UI regardless of manifest")
    args = ap.parse_args(argv)
    summary = run(args.before, args.after, args.features, args.typed40, args.ui)
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as fh:
        json.dump(summary, fh, indent=2, default=str)
    for feat, e in summary["features"].items():
        extra = ""
        if "thresholds" in e and e["thresholds"]["failures"]:
            extra = " | " + "; ".join(e["thresholds"]["failures"])
        print(f"{feat}: {'PASS' if e['passed'] else 'FAIL'} ({e.get('classification')}){extra}")
    print(f"OVERALL: {'PASS' if summary['passed'] else 'FAIL'} -> {args.out}")
    return 0 if summary["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
