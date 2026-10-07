"""Unit tests for tools/compare_grounding_regression.py (synthetic workbooks only).

Spec: tsg-nbop-crawl-grounded-ui-steps, task 13.1 (Req 8.1, 8.2, 8.4, 9.1-9.6).
"""
import importlib.util
import json
import os

import openpyxl
import pytest

_TOOL = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                     "tools", "compare_grounding_regression.py")
_spec = importlib.util.spec_from_file_location("compare_grounding_regression", _TOOL)
cgr = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(cgr)

TC_HEADER = ["S.No", "Summary", "Description", "Step #", "Step Summary", "Expected Result"]
CONTRACT_HEADER = ["TC Number", "Step Number", "Action ID", "Parameters"]


def _tc_rows(titles, step_text="Do the thing"):
    rows = []
    for i, t in enumerate(titles, 1):
        rows.append([str(i), t, "desc", 1, step_text, "ok"])
        rows.append([None, None, None, 2, "Second step", "ok"])
    return rows


def _write(path, titles, step_text="Do the thing", contract=None, report=None, extra_sheet=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    ws.append(TC_HEADER)
    for r in _tc_rows(titles, step_text):
        ws.append(r)
    wb.create_sheet("Summary").append(["x"])
    cs = wb.create_sheet("Execution Contract")
    cs.append(CONTRACT_HEADER)
    for r in contract or [["TC01", 1, "legacy", "{}"]]:
        cs.append(r)
    if extra_sheet:
        wb.create_sheet(extra_sheet).append(["y"])
    if report is not None:
        metrics, issues = report
        rs = wb.create_sheet("Grounding Report")
        for label, value in metrics:
            rs.append([label, value])
        rs.append([])
        rs.append(["TC Number", "Step Number", "Step UID", "Step Text", "Reason", "Candidate Pages"])
        for r in issues:
            rs.append(r)
    wb.save(path)
    return str(path)


def _metrics(ui=10, grounded=7, unmatched=3, filler=0, typed=5, total=10,
             gshare="70.0%", tshare="50.0%"):
    return [
        ("Feature", "SYN-1"), ("Crawl Map", "map.json"), ("Crawl discovered_at", "2026-04-14"),
        ("Crawl age (days)", 3), ("Staleness warning", None),
        ("UI_Steps", ui), ("Grounded_Steps", grounded), ("Unmatched_Steps", unmatched),
        ("Filler_Steps (post-pass)", filler), ("Closer_Steps removed", 2),
        ("Typed rows", typed), ("Total rows", total),
        ("Grounded share %", gshare), ("Typed share %", tshare),
    ]


def _issues(n_unmatched=3, extra_unsupported=1):
    rows = [[f"TC0{i+1}", 1, f"u{i}", "text", "no_match", ""] for i in range(n_unmatched)]
    rows += [["TC09", 2, "ux", "text", "action_unsupported", "P"]] * extra_unsupported
    return rows


# --- compare_feature --------------------------------------------------------

def test_equal_workbooks_pass_non_ui(tmp_path):
    b = _write(tmp_path / "b.xlsx", ["TC01_A", "TC02_B"])
    a = _write(tmp_path / "a.xlsx", ["TC01_A", "TC02_B"])
    r = cgr.compare_feature(b, a, is_ui=False)
    assert r["passed"] and r["tc_count_before"] == 2
    assert r["test_cases_identical"] and r["contract_identical"] and r["sheet_names_equal"]


def test_title_change_fails(tmp_path):
    b = _write(tmp_path / "b.xlsx", ["TC01_A", "TC02_B"])
    a = _write(tmp_path / "a.xlsx", ["TC01_A", "TC02_CHANGED"])
    r = cgr.compare_feature(b, a, is_ui=True)
    assert r["tc_count_equal"] and not r["titles_equal"] and not r["passed"]
    assert r["title_diffs"] == [{"index": 2, "before": "TC02_B", "after": "TC02_CHANGED"}]


def test_tc_count_change_fails(tmp_path):
    b = _write(tmp_path / "b.xlsx", ["TC01_A", "TC02_B"])
    a = _write(tmp_path / "a.xlsx", ["TC01_A"])
    r = cgr.compare_feature(b, a, is_ui=True)
    assert not r["tc_count_equal"] and not r["passed"]


def test_non_ui_step_cell_diff_fails_but_ui_tolerates(tmp_path):
    b = _write(tmp_path / "b.xlsx", ["TC01_A"], step_text="Verify requirement: X")
    a = _write(tmp_path / "a.xlsx", ["TC01_A"], step_text="On NBOP 'Home', verify 'X' is displayed")
    r = cgr.compare_feature(b, a, is_ui=False)
    assert not r["test_cases_identical"] and not r["passed"]
    assert r["test_cases_first_diff"]["row"] == 2
    assert cgr.compare_feature(b, a, is_ui=True)["passed"]


def test_non_ui_contract_diff_and_sheet_names(tmp_path):
    b = _write(tmp_path / "b.xlsx", ["TC01_A"])
    a = _write(tmp_path / "a.xlsx", ["TC01_A"], contract=[["TC01", 1, "nbop.ui.click", "{}"]])
    r = cgr.compare_feature(b, a, is_ui=False)
    assert not r["contract_identical"] and not r["passed"]
    c = _write(tmp_path / "c.xlsx", ["TC01_A"], extra_sheet="Grounding Report")
    r2 = cgr.compare_feature(b, c, is_ui=False)
    assert not r2["sheet_names_equal"] and not r2["passed"]


# --- ui_metrics + thresholds -----------------------------------------------

def test_report_meeting_all_thresholds(tmp_path):
    a = _write(tmp_path / "a.xlsx", ["TC01_A"], report=(_metrics(), _issues()))
    m = cgr.ui_metrics(a)
    assert m["filler_steps"] == 0 and m["grounded_share"] == 70.0 and m["typed_share"] == 50.0
    assert m["unmatched_issue_rows"] == 3 and m["issue_rows"] == 4
    assert m["issues_match_unmatched"] and m["all_ui_classified"]
    assert len(m["no_match_rows"]) == 3
    assert cgr.check_thresholds(m, typed40=True)["passed"]


@pytest.mark.parametrize("kwargs,typed40,needle", [
    ({"filler": 1}, False, "filler_steps"),
    ({"gshare": "54.9%"}, False, "grounded_share"),
    ({"tshare": "39.0%"}, True, "typed_share"),
    ({"tshare": 14.0}, False, "typed_share"),
    ({"grounded": 6}, False, "grounded + unmatched"),
])
def test_report_missing_each_threshold(tmp_path, kwargs, typed40, needle):
    a = _write(tmp_path / "a.xlsx", ["TC01_A"], report=(_metrics(**kwargs), _issues()))
    th = cgr.check_thresholds(cgr.ui_metrics(a), typed40=typed40)
    assert not th["passed"] and any(needle in f for f in th["failures"])


def test_typed_share_15_ok_without_typed40(tmp_path):
    a = _write(tmp_path / "a.xlsx", ["TC01_A"], report=(_metrics(tshare="20%"), _issues()))
    assert cgr.check_thresholds(cgr.ui_metrics(a), typed40=False)["passed"]


def test_issue_count_mismatch_fails(tmp_path):
    a = _write(tmp_path / "a.xlsx", ["TC01_A"], report=(_metrics(), _issues(n_unmatched=2)))
    th = cgr.check_thresholds(cgr.ui_metrics(a), typed40=False)
    assert not th["passed"] and any("Unmatched_Steps" in f for f in th["failures"])


def test_missing_report_sheet(tmp_path):
    a = _write(tmp_path / "a.xlsx", ["TC01_A"])
    m = cgr.ui_metrics(a)
    assert m == {"has_grounding_report": False}
    assert not cgr.check_thresholds(m, typed40=False)["passed"]


# --- CLI --------------------------------------------------------------------

def test_cli_uses_manifest_and_writes_summary(tmp_path):
    before, after = tmp_path / "before", tmp_path / "after"
    before.mkdir(); after.mkdir()
    _write(before / "SYN-1.xlsx", ["TC01_A"])
    _write(after / "SYN-1.xlsx", ["TC01_A"])
    _write(before / "SYN-2.xlsx", ["TC01_U"])
    _write(after / "SYN-2.xlsx", ["TC01_U"], report=(_metrics(tshare="30%"), _issues()))
    (before / "manifest.json").write_text(json.dumps([
        {"feature": "SYN-1", "command": "x", "source_path": "x", "classification": "api"},
        {"feature": "SYN-2", "command": "x", "source_path": "x", "classification": "ui"},
    ]), encoding="utf-8")
    out = tmp_path / "summary.json"
    rc = cgr.main(["--before", str(before), "--after", str(after), "--out", str(out)])
    s = json.loads(out.read_text(encoding="utf-8"))
    assert rc == 0 and s["passed"]
    assert s["features"]["SYN-1"]["comparison"]["test_cases_identical"]
    assert "metrics" not in s["features"]["SYN-1"]
    assert s["features"]["SYN-2"]["thresholds"]["typed_min"] == 15.0

    rc2 = cgr.main(["--before", str(before), "--after", str(after), "--out", str(out),
                    "--features", "SYN-2", "--typed40", "SYN-2"])
    s2 = json.loads(out.read_text(encoding="utf-8"))
    assert rc2 == 1 and not s2["passed"] and list(s2["features"]) == ["SYN-2"]
