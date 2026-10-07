# -*- coding: utf-8 -*-
"""Task 9.2 — contract-writer branch in ``_build_execution_contract_sheets``.

Validates: Requirements 3.4, 4.1, 4.2, 4.3, 4.4, 4.5, 4.6

These tests build a tiny suite in memory, attach ``_crawl_grounding`` records to
the steps of one UI feature TC (typed / grounded-unregistered / unmatched cases)
plus a plain non-UI TC, run ``_build_execution_contract_sheets`` against a fresh
workbook, and read the "Execution Contract" sheet back.

They assert:
  * UI rows use the UI action ids / typed mode and the crawl provenance;
  * grounded-but-unregistered rows stay legacy but keep the TC score and crawl
    provenance;
  * unmatched rows fall back to the generic provenance and 0% grounding;
  * the non-UI rows are byte-identical to a run with no grounding attached;
  * ``suite._grounding_result.typed_rows`` / ``total_rows`` are set from the UI
    branch only.

No Jira, Chalk, cache or browser is touched.
"""
import os
import sys
from types import SimpleNamespace

import openpyxl
import pytest

_REPO_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if _REPO_ROOT not in sys.path:
    sys.path.insert(0, _REPO_ROOT)
_TSG_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _TSG_ROOT not in sys.path:
    sys.path.insert(0, _TSG_ROOT)

import modules.excel_generator as eg  # noqa: E402
from modules.contract_bridge import CONTRACT_HEADERS, CONTRACT_SHEET  # noqa: E402
from modules.crawl_grounding import GroundingResult  # noqa: E402
from modules.data_models_v8 import (  # noqa: E402
    RoutingAudit, TestCase, TestStep, TestSuite,
)

_LOCATOR = ("TMO DashBoard/nbop_discovery/nbop_ui_map.json"
            "#pages['Subscriber Profile'].tabs['Usage Details']")


def _typed_grounding():
    """A grounded step whose action id is registered -> typed ui_action row."""
    return SimpleNamespace(
        kind="ui_action", reason="", page="Subscriber Profile",
        page_url="https://nbop.example.com/subscriber/profile", page_kind="page",
        control_type="tab", control_label="Usage Details",
        action_id="nbop.ui.open_tab", value="", expect="",
        grounding_percent=100, source_locator=_LOCATOR,
    )


def _unregistered_grounding():
    """Grounded but the action id has no TSE handler -> legacy, keeps TC score."""
    return SimpleNamespace(
        kind="unmatched", reason="action_unsupported", page="Subscriber Profile",
        page_url="https://nbop.example.com/subscriber/profile", page_kind="page",
        control_type="input", control_label="MSISDN",
        action_id="nbop.ui.enter_text", value="5551234567", expect="",
        grounding_percent=None, source_locator=_LOCATOR,
    )


def _unmatched_grounding():
    """No crawl match -> legacy, 0% grounding, keep existing provenance."""
    return SimpleNamespace(
        kind="unmatched", reason="no_match", page="", page_url="", page_kind="",
        control_type="", control_label="", action_id="", value="", expect="",
        grounding_percent=0, source_locator="",
    )


def _ui_tc():
    """A UI TC: three steps, each with a prepared ``_crawl_grounding`` record."""
    s_typed = TestStep(step_num=1, summary="Click the Usage Details tab",
                       expected="Usage Details tab opens")
    s_typed._crawl_grounding = _typed_grounding()
    s_unreg = TestStep(step_num=2, summary="Enter MSISDN 5551234567",
                       expected="Value entered")
    s_unreg._crawl_grounding = _unregistered_grounding()
    s_unmatched = TestStep(step_num=3, summary="[UNMATCHED-CRAWL] Verify the gizmo",
                           expected="")
    s_unmatched._crawl_grounding = _unmatched_grounding()
    return TestCase(sno="1", summary="UI flow", steps=[s_typed, s_unreg, s_unmatched],
                    grounding_score=42)


def _nonui_tc():
    """A plain TC with no grounding attached — must stay on the generic path."""
    step = TestStep(step_num=1, summary="Call the provisioning API",
                    expected="HTTP 200")
    return TestCase(sno="2", summary="API flow", steps=[step], grounding_score=77)


def _rows_by_tc(ws):
    """Map Execution Contract rows -> list of dicts keyed by header, per TC number."""
    header = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(header, row))
        out.setdefault(str(rec["TC Number"]), []).append(rec)
    return out


def _build(suite):
    wb = openpyxl.Workbook()
    eg._build_execution_contract_sheets(wb, suite, log=lambda *_: None)
    return wb[CONTRACT_SHEET]


def test_ui_rows_use_ui_contract_and_nonui_unchanged():
    # Suite WITH grounding on the UI TC.
    grounded_suite = TestSuite(
        feature_id="MWTGPROV-9002", feature_title="UI Branch Fixture", pi="PI-59",
        test_cases=[_ui_tc(), _nonui_tc()],
        routing_audit=RoutingAudit(classification="ui"),
    )
    grounded_suite._grounding_result = GroundingResult(
        feature_id="MWTGPROV-9002", crawl_path="", crawl_error="",
        discovered_at="", age_days=None, stale_warning="",
        ui_steps=3, grounded_steps=2, unmatched_steps=1,
        filler_steps=0, closers_removed=0,
    )
    ws = _build(grounded_suite)

    # Headers are unchanged (Req 4.6).
    assert [c.value for c in ws[1]] == CONTRACT_HEADERS

    by_tc = _rows_by_tc(ws)
    ui_rows = by_tc["1"]
    assert len(ui_rows) == 3

    # Row 1: typed ui_action, crawl provenance, 100%.
    typed_row = ui_rows[0]
    assert typed_row["Mode"] == "typed"
    assert typed_row["Action Type"] == "ui_action"
    assert typed_row["Action ID"] == "nbop.ui.open_tab"
    assert typed_row["Target System"] == "NBOP"
    assert typed_row["Source Type"] == "nbop_crawl"
    assert typed_row["Source Locator"] == _LOCATOR
    assert int(typed_row["Grounding %"]) == 100

    # Row 2: grounded but unregistered -> legacy, keeps TC score (42) and crawl
    # provenance; the id survives for diagnostics (Req 4.4).
    unreg_row = ui_rows[1]
    assert unreg_row["Mode"] == "legacy"
    assert unreg_row["Action Type"] == "legacy"
    assert unreg_row["Action ID"] == "nbop.ui.enter_text"
    assert unreg_row["Source Type"] == "nbop_crawl"
    assert int(unreg_row["Grounding %"]) == 42  # TC score preserved, not None/0

    # Row 3: unmatched -> legacy, 0%, generic provenance preserved (not nbop_crawl).
    unmatched_row = ui_rows[2]
    assert unmatched_row["Mode"] == "legacy"
    assert unmatched_row["Action ID"] == ""
    assert int(unmatched_row["Grounding %"]) == 0
    assert unmatched_row["Source Type"] != "nbop_crawl"

    # Counts come from the UI branch only (3 UI steps, 1 typed).
    assert grounded_suite._grounding_result.typed_rows == 1
    assert grounded_suite._grounding_result.total_rows == 3

    # Non-UI TC: compare against a suite built WITHOUT any grounding records.
    plain_suite = TestSuite(
        feature_id="MWTGPROV-9002", feature_title="UI Branch Fixture", pi="PI-59",
        test_cases=[_ui_tc_plain(), _nonui_tc()],
        routing_audit=RoutingAudit(classification="ui"),
    )
    plain_ws = _build(plain_suite)
    plain_by_tc = _rows_by_tc(plain_ws)
    assert by_tc["2"] == plain_by_tc["2"], "non-UI rows must be unchanged by the UI branch"


def _ui_tc_plain():
    """Same shape as ``_ui_tc`` but with NO ``_crawl_grounding`` attached."""
    s1 = TestStep(step_num=1, summary="Click the Usage Details tab",
                  expected="Usage Details tab opens")
    s2 = TestStep(step_num=2, summary="Enter MSISDN 5551234567", expected="Value entered")
    s3 = TestStep(step_num=3, summary="[UNMATCHED-CRAWL] Verify the gizmo", expected="")
    return TestCase(sno="1", summary="UI flow", steps=[s1, s2, s3], grounding_score=42)


def test_grounding_result_untouched_when_absent():
    """A suite without ``_grounding_result`` must not gain one, and no crash."""
    suite = TestSuite(
        feature_id="MWTGPROV-9003", feature_title="No Result Fixture", pi="PI-59",
        test_cases=[_nonui_tc()],
    )
    ws = _build(suite)
    assert [c.value for c in ws[1]] == CONTRACT_HEADERS
    assert getattr(suite, "_grounding_result", None) is None


def test_steps_without_grounding_use_generic_path():
    """A UI-classified suite whose steps carry no grounding stays byte-identical."""
    with_result = TestSuite(
        feature_id="MWTGPROV-9004", feature_title="Plain UI", pi="PI-59",
        test_cases=[_ui_tc_plain()],
        routing_audit=RoutingAudit(classification="ui"),
    )
    with_result._grounding_result = GroundingResult(
        feature_id="MWTGPROV-9004", crawl_path="", crawl_error="",
        discovered_at="", age_days=None, stale_warning="",
        ui_steps=0, grounded_steps=0, unmatched_steps=0,
        filler_steps=0, closers_removed=0,
    )
    ws = _build(with_result)
    rows = _rows_by_tc(ws)["1"]
    # No step had _crawl_grounding, so no UI rows were counted.
    assert with_result._grounding_result.typed_rows == 0
    assert with_result._grounding_result.total_rows == 0
    # All three generic rows are legacy (none match infer_step_contract typed ids).
    assert all(r["Mode"] == "legacy" for r in rows)
