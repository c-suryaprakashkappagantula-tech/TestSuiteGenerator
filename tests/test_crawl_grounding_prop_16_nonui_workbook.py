# -*- coding: utf-8 -*-
"""Property 16 (workbook part): Non-UI suites are byte-identical at the workbook level.

A suite whose ``routing_audit.classification`` is not a UI_Feature (``api``, or
``hybrid`` carrying an API signal) must produce the *same workbook* whether the
crawl-grounding pass runs for real (where it is a no-op early-return for non-UI
suites) or is stubbed out to a pure no-op. Requirement 8.2 says Non-UI
``Test Cases`` and ``Execution Contract`` rows are identical to the pre-change
output; this test pins that the pass itself cannot perturb a non-UI workbook.

For each generated non-UI suite ``generate_excel`` is run twice into a tmp dir:

  (a) once with the real :func:`modules.crawl_grounding.apply_crawl_grounding`
      (a no-op for non-UI suites), and
  (b) once with ``apply_crawl_grounding`` monkeypatched to ``lambda *a, **k: None``.

Equality strategy: openpyxl writes a zip whose entries carry timestamps, so raw
``.xlsx`` bytes are not stable across two writes. We therefore compare at the
cell level — identical sheet names, and for every sheet identical cell values in
every row/column. This is the meaningful, stable equality for "same workbook".

Both runs must also leave the non-UI suite with NO ``_grounding_result`` and the
workbook with NO ``Grounding Report`` sheet, and neither run may emit a
``[CRAWL-GROUND] ... ui_steps=`` summary line.

Hermetic: ``output_path``/``checkpoint_path`` are redirected into a tmp dir so
nothing is written to the real outputs directory, and crawl grounding never
reaches the real ``nbop_ui_map.json`` (the real pass early-returns for non-UI
suites before loading any map; the stubbed pass does nothing). A fresh
``MonkeyPatch`` and ``TemporaryDirectory`` are created per example rather than
using pytest's function-scoped fixtures, which hypothesis would reject inside
``@given``.

# Feature: tsg-nbop-crawl-grounded-ui-steps, Property 16: Non-UI suites are untouched (workbook part)
Validates: Requirements 1.2, 6.5, 8.2
"""
import copy
import os
import sys
import tempfile
from pathlib import Path

import openpyxl
import pytest
from hypothesis import given, settings

_TESTS_DIR = os.path.dirname(os.path.abspath(__file__))
_TSG_ROOT = os.path.dirname(_TESTS_DIR)
_REPO_ROOT = os.path.dirname(_TSG_ROOT)
for _p in (_TESTS_DIR, _TSG_ROOT, _REPO_ROOT):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import modules.excel_generator as eg  # noqa: E402
import modules.crawl_grounding as cg  # noqa: E402
from crawl_grounding_strategies import suites  # noqa: E402


def _sheet_cells(path):
    """Return (sheet_names, {sheet_name: tuple-of-row-tuples}) for a workbook.

    Cell-level snapshot used for equality: sheet order/names plus every cell
    value. Stable across repeated writes, unlike the raw zip bytes (which carry
    per-entry timestamps).
    """
    wb = openpyxl.load_workbook(str(path), data_only=False)
    try:
        snap = {}
        for name in wb.sheetnames:
            ws = wb[name]
            snap[name] = tuple(
                tuple(cell.value for cell in row)
                for row in ws.iter_rows()
            )
        return list(wb.sheetnames), snap
    finally:
        wb.close()


def _run_generate(suite, out_file, cp_file, mp, stub_grounding):
    """Run ``generate_excel`` into fixed paths, capturing the log lines.

    When ``stub_grounding`` is True the grounding pass is replaced by a pure
    no-op; otherwise the real pass runs (a no-op early-return for non-UI
    suites). Returns the list of log messages emitted during the run.
    """
    mp.setattr(eg, 'output_path', lambda *a, **k: out_file)
    mp.setattr(eg, 'checkpoint_path', lambda *a, **k: cp_file)
    if stub_grounding:
        mp.setattr(cg, 'apply_crawl_grounding', lambda *a, **k: None)

    logs = []
    result = eg.generate_excel(suite, log=logs.append)
    assert result == out_file
    assert out_file.exists()
    return logs


def _assert_no_grounding(suite, logs, sheet_names):
    """Non-UI invariants: no grounding artefacts, no report sheet, no summary log."""
    assert not hasattr(suite, '_grounding_result'), \
        'non-UI suite must not gain a _grounding_result'
    for tc in getattr(suite, 'test_cases', None) or []:
        for step in getattr(tc, 'steps', None) or []:
            assert not hasattr(step, '_crawl_grounding'), \
                'non-UI step must not gain a _crawl_grounding'
    assert 'Grounding Report' not in sheet_names, \
        'non-UI workbook must have no Grounding Report sheet'
    assert not any('[CRAWL-GROUND]' in str(m) and 'ui_steps=' in str(m) for m in logs), \
        'non-UI run must not log a [CRAWL-GROUND] ... ui_steps= summary line'


@settings(max_examples=30, deadline=None)
@given(suite=suites(False))
def test_non_ui_workbook_identical_with_and_without_pass(suite):
    # A fresh MonkeyPatch + tmp dir per example (hypothesis rejects pytest's
    # function-scoped tmp_path/monkeypatch fixtures inside @given).
    with pytest.MonkeyPatch.context() as mp, tempfile.TemporaryDirectory() as td:
        tmp = Path(td)
        # Two independent deepcopies: generate_excel mutates the suite
        # (sanitize/polish/dedupe) in place, so each run gets its own copy.
        suite_real = copy.deepcopy(suite)
        suite_stub = copy.deepcopy(suite)

        out_real = tmp / 'suite_real.xlsx'
        cp_real = tmp / 'suite_real_checkpoint.xlsx'
        out_stub = tmp / 'suite_stub.xlsx'
        cp_stub = tmp / 'suite_stub_checkpoint.xlsx'

        # (a) real grounding pass (no-op for non-UI).
        logs_real = _run_generate(suite_real, out_real, cp_real, mp, stub_grounding=False)
        names_real, cells_real = _sheet_cells(out_real)
        _assert_no_grounding(suite_real, logs_real, names_real)

        # (b) grounding pass stubbed to a pure no-op.
        logs_stub = _run_generate(suite_stub, out_stub, cp_stub, mp, stub_grounding=True)
        names_stub, cells_stub = _sheet_cells(out_stub)
        _assert_no_grounding(suite_stub, logs_stub, names_stub)

        # Identical workbooks: same sheet names in the same order, and identical
        # cell values on every sheet (Test Cases, Execution Contract, the rest).
        assert names_real == names_stub, (
            f'sheet names differ:\n  real={names_real!r}\n  stub={names_stub!r}'
        )
        for name in names_real:
            assert cells_real[name] == cells_stub[name], (
                f'cell values differ on sheet {name!r} between the real-pass and '
                f'no-op-pass workbooks'
            )
